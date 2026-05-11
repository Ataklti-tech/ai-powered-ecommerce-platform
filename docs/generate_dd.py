from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

# ─── Helpers ──────────────────────────────────────────────────────────────────

def set_cell_border(cell, top=True, bottom=True, left=True, right=True):
    """Add thin black borders to a cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    sides = {}
    if top:    sides['top']    = borders
    if bottom: sides['bottom'] = borders
    if left:   sides['left']   = borders
    if right:  sides['right']  = borders

    for side in ['top', 'bottom', 'left', 'right']:
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'),   'single')
        el.set(qn('w:sz'),    '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    tcPr.append(borders)

def set_row_shading(row, shade=False):
    for cell in row.cells:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        if shade:
            shd.set(qn('w:val'),   'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'),  'E8E8E8')   # light grey header
        else:
            shd.set(qn('w:val'),   'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'),  'FFFFFF')
        tcPr.append(shd)

def style_table(table):
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell)
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(2)
                para.paragraph_format.space_after  = Pt(2)

def add_header_row(table, headers, col_widths=None):
    row = table.rows[0]
    set_row_shading(row, shade=True)
    for i, text in enumerate(headers):
        cell = row.cells[i]
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = para.add_run(text)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0, 0, 0)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

def add_data_row(table, values, bold_first=False):
    row = table.add_row()
    set_row_shading(row, shade=False)
    for i, text in enumerate(values):
        cell = row.cells[i]
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = para.add_run(str(text))
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0, 0, 0)
        if bold_first and i == 0:
            run.bold = True
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

def set_col_widths(table, widths_inches):
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths_inches):
                cell.width = Inches(widths_inches[i])

def add_section_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16) if level == 1 else Pt(10)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12) if level == 1 else Pt(10)
    run.font.color.rgb = RGBColor(0, 0, 0)

def add_caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(6)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(80, 80, 80)

# ─── Document setup ───────────────────────────────────────────────────────────

def new_doc(title):
    doc = Document()
    # Margins
    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1.1)
        section.right_margin  = Inches(1.1)

    # Title
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0, 0, 0)

    # Subtitle line
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(20)
    run2 = p2.add_run('E-Commerce Platform  |  System Documentation')
    run2.font.size = Pt(9)
    run2.font.color.rgb = RGBColor(80, 80, 80)

    return doc

# ═══════════════════════════════════════════════════════════════════════════════
# DOCUMENT 1 — DATA DICTIONARY: ERD
# ═══════════════════════════════════════════════════════════════════════════════

doc1 = new_doc('Data Dictionary — Entity Relationship Diagram')

entities = [
    {
        'name': 'User',
        'desc': 'Stores all registered user accounts, credentials, and embedded shipping addresses.',
        'fields': [
            ('_id',                   'ObjectId', 'PK',               'Auto-generated MongoDB document identifier'),
            ('name',                  'String',   'Required',          'Full display name of the user'),
            ('email',                 'String',   'Required, Unique',  'Login email address'),
            ('password',              'String',   'Required',          'bcrypt-hashed password (never stored in plain text)'),
            ('role',                  'String',   'Required',          'Enum: user | admin'),
            ('isVerified',            'Boolean',  'Optional',          'True once the user confirms their email'),
            ('addresses',             'Array',    'Optional',          'Embedded address objects: {street, city, postalCode, country, isDefault}'),
            ('passwordResetToken',    'String',   'Optional',          'SHA-256 hashed token for password reset flow'),
            ('passwordResetExpires',  'Date',     'Optional',          'Expiry timestamp for the reset token'),
            ('createdAt',             'Date',     'Auto',              'Account creation timestamp (Mongoose timestamps)'),
            ('updatedAt',             'Date',     'Auto',              'Last modification timestamp (Mongoose timestamps)'),
        ]
    },
    {
        'name': 'Product',
        'desc': 'Catalogue of all products available for purchase.',
        'fields': [
            ('_id',             'ObjectId',      'PK',                'Auto-generated document identifier'),
            ('name',            'String',        'Required',          'Product display name'),
            ('slug',            'String',        'Unique',            'URL-friendly identifier, auto-derived from name'),
            ('description',     'String',        'Required',          'Full product description'),
            ('shortDescription','String',        'Optional',          'Brief summary shown on product cards'),
            ('price',           'Number',        'Required',          'Current selling price in base currency'),
            ('comparePrice',    'Number',        'Optional',          'Original price shown as strikethrough'),
            ('stock',           'Number',        'Required',          'Available inventory count'),
            ('status',          'String',        'Required',          'Enum: active | inactive | draft'),
            ('category',        'ObjectId',      'FK → Category',     'Primary product category reference'),
            ('images',          'Array',         'Required',          'Cloudinary image objects: [{url, publicId}]'),
            ('averageRating',   'Number',        'Optional',          'Computed average of all review ratings (0–5)'),
            ('reviewCount',     'Number',        'Optional',          'Total number of reviews for this product'),
            ('tags',            'Array<String>', 'Optional',          'Keywords for search and filtering'),
        ]
    },
    {
        'name': 'Category',
        'desc': 'Hierarchical product categories. A category may have a parent for sub-categories.',
        'fields': [
            ('_id',         'ObjectId', 'PK',                     'Auto-generated document identifier'),
            ('name',        'String',   'Required, Unique',        'Category display name'),
            ('slug',        'String',   'Unique',                  'URL-friendly identifier'),
            ('description', 'String',   'Optional',                'Category description shown on listing pages'),
            ('image',       'String',   'Optional',                'Banner or icon URL for the category'),
            ('parent',      'ObjectId', 'FK → Category, Optional', 'Reference to parent category (enables hierarchy)'),
        ]
    },
    {
        'name': 'Order',
        'desc': 'Records a customer purchase from placement through fulfilment.',
        'fields': [
            ('_id',             'ObjectId', 'PK',             'Auto-generated document identifier'),
            ('user',            'ObjectId', 'FK → User, Req', 'Reference to the ordering customer'),
            ('orderNumber',     'String',   'Unique',          'Human-readable reference (e.g. ORD-00045)'),
            ('items',           'Array',    'Required',        'Line items: [{product, name, price, qty, image}]'),
            ('shippingAddress', 'Object',   'Required',        'Address snapshot captured at checkout time'),
            ('status',          'String',   'Required',        'Enum: pending | confirmed | processing | shipped | delivered | cancelled | refunded'),
            ('paymentMethod',   'String',   'Required',        'Enum: card | mtn_momo | airtel | pesapal | paypal | stripe'),
            ('paymentStatus',   'String',   'Required',        'Enum: pending | paid | failed | refunded'),
            ('subtotal',        'Number',   'Required',        'Sum of item prices before discount and shipping'),
            ('discount',        'Number',   'Optional',        'Discount amount from applied coupon'),
            ('shippingCost',    'Number',   'Optional',        'Shipping fee added to the total'),
            ('total',           'Number',   'Required',        'Final amount charged to the customer'),
            ('couponCode',      'String',   'Optional',        'Coupon code applied at checkout'),
        ]
    },
    {
        'name': 'Cart',
        'desc': 'Persisted shopping cart — one cart per registered user.',
        'fields': [
            ('_id',        'ObjectId', 'PK',                   'Auto-generated document identifier'),
            ('user',       'ObjectId', 'FK → User, Unique',    'Cart owner; unique index ensures one cart per user'),
            ('items',      'Array',    'Optional',              'Cart line items: [{product, name, price, qty, image}]'),
            ('total',      'Number',   'Optional',              'Computed sum of all item prices'),
            ('couponCode', 'String',   'Optional',              'Coupon code applied to the cart'),
            ('discount',   'Number',   'Optional',              'Discount amount from the applied coupon'),
        ]
    },
    {
        'name': 'Payment',
        'desc': 'Transaction record for every payment attempt across all gateways.',
        'fields': [
            ('_id',             'ObjectId', 'PK',                 'Auto-generated document identifier'),
            ('order',           'ObjectId', 'FK → Order, Req',    'Order this payment is associated with'),
            ('user',            'ObjectId', 'FK → User, Req',     'Customer who made the payment'),
            ('method',          'String',   'Required',            'Enum: stripe | paypal | mtn_momo | airtel | pesapal'),
            ('gateway',         'String',   'Optional',            'Actual gateway used (e.g. pesapal aggregating MTN MoMo)'),
            ('status',          'String',   'Required',            'Enum: pending | processing | completed | failed | refunded'),
            ('amount',          'Number',   'Required',            'Amount charged'),
            ('currency',        'String',   'Required',            'ISO 4217 code: UGX, USD, etc.'),
            ('transactionId',   'String',   'Optional',            'Gateway-issued transaction reference'),
            ('gatewayResponse', 'Mixed',    'Optional',            'Raw webhook or callback payload stored for audit'),
        ]
    },
    {
        'name': 'Review',
        'desc': 'Customer reviews and star ratings attached to products.',
        'fields': [
            ('_id',             'ObjectId', 'PK',                   'Auto-generated document identifier'),
            ('user',            'ObjectId', 'FK → User, Required',  'Review author'),
            ('product',         'ObjectId', 'FK → Product, Req',    'Reviewed product'),
            ('rating',          'Number',   'Required',              'Star rating 1–5'),
            ('title',           'String',   'Optional',              'Review headline'),
            ('body',            'String',   'Required',              'Full review text'),
            ('verifiedPurchase','Boolean',  'Optional',              'True if reviewer has a completed order for this product'),
            ('createdAt',       'Date',     'Auto',                  'Submission timestamp'),
        ]
    },
    {
        'name': 'Wishlist',
        'desc': 'Products saved by a user for later viewing — one wishlist per user.',
        'fields': [
            ('_id',   'ObjectId', 'PK',                  'Auto-generated document identifier'),
            ('user',  'ObjectId', 'FK → User, Unique',   'Wishlist owner; unique index ensures one per user'),
            ('items', 'Array',    'Optional',             'Saved products: [{product (ObjectId), addedAt (Date)}]'),
        ]
    },
    {
        'name': 'UserActivity',
        'desc': 'Behavioural event log used for analytics and AI recommendation model training.',
        'fields': [
            ('_id',          'ObjectId', 'PK',                    'Auto-generated document identifier'),
            ('user',         'ObjectId', 'FK → User, Required',   'User who performed the action'),
            ('product',      'ObjectId', 'FK → Product, Optional','Product involved in the event'),
            ('activityType', 'String',   'Required',               'Enum: view | search | cart_add | cart_remove | wishlist_add | purchase | review | rec_impression | similar_impression'),
            ('metadata',     'Mixed',    'Optional',               'Free-form context: search query, referrer URL, recommendation type, etc.'),
            ('timestamp',    'Date',     'Required',               'Event time; indexed for range queries and analytics'),
        ]
    },
    {
        'name': 'Coupon',
        'desc': 'Discount codes that can be applied at cart or checkout.',
        'fields': [
            ('_id',            'ObjectId', 'PK',                 'Auto-generated document identifier'),
            ('code',           'String',   'Required, Unique',   'Coupon code string (e.g. SAVE20)'),
            ('discountType',   'String',   'Required',           'Enum: percentage | fixed'),
            ('discountValue',  'Number',   'Required',           'Discount amount or percentage value'),
            ('minOrderAmount', 'Number',   'Optional',           'Minimum cart total required to apply the coupon'),
            ('usageLimit',     'Number',   'Optional',           'Maximum number of total redemptions allowed'),
            ('usageCount',     'Number',   'Optional',           'Running count of how many times the coupon has been used'),
            ('isActive',       'Boolean',  'Required',           'Whether the coupon is enabled'),
            ('expiresAt',      'Date',     'Optional',           'Coupon expiry date and time'),
        ]
    },
]

HEADERS = ['Field', 'Type', 'Constraint', 'Description']
WIDTHS  = [1.2, 1.0, 1.3, 3.3]   # total ~6.8"

for entity in entities:
    add_section_heading(doc1, entity['name'], level=1)
    add_caption(doc1, entity['desc'])

    table = doc1.add_table(rows=1, cols=4)
    style_table(table)
    add_header_row(table, HEADERS)
    for row_data in entity['fields']:
        add_data_row(table, row_data, bold_first=True)
    set_col_widths(table, WIDTHS)

    doc1.add_paragraph()  # spacing after table

doc1.save(r'c:\Users\hp\Desktop\e_commerce\docs\DD_ERD.docx')
print('DD_ERD.docx saved')

# ═══════════════════════════════════════════════════════════════════════════════
# DOCUMENT 2 — DATA DICTIONARY: DFD
# ═══════════════════════════════════════════════════════════════════════════════

doc2 = new_doc('Data Dictionary — Data Flow Diagram (Level 1)')

# ── External Entities ─────────────────────────────────────────────────────────
add_section_heading(doc2, '1. External Entities', level=1)
add_caption(doc2, 'Actors and systems that interact with the e-commerce platform from outside its boundary.')

ext_headers = ['ID', 'Name', 'Description', 'Inputs to System', 'Outputs from System']
ext_widths  = [0.35, 1.2, 1.8, 1.8, 1.65]

ext_data = [
    ('E1', 'Customer',         'End user who browses, purchases, and reviews products.',
     'Credentials, checkout data, reviews, activity events',
     'Products, order status, recommendations, receipts'),
    ('E2', 'Admin',            'Staff member managing inventory, orders, and platform settings.',
     'Product data, category updates, order status changes',
     'Analytics reports, dashboard metrics, alerts'),
    ('E3', 'Payment Gateway',  'External payment processors: Pesapal, MTN MoMo, Airtel Money, Stripe, PayPal.',
     'Transaction result, IPN / webhook callback',
     'Payment initiation request (amount, reference, redirect URLs)'),
    ('E4', 'AI Service',       'FastAPI microservice hosting collaborative-filter and content-based recommendation models.',
     'Ranked recommendation IDs, similarity scores',
     'User behaviour summary, product ID lists for scoring'),
]

t_ext = doc2.add_table(rows=1, cols=5)
style_table(t_ext)
add_header_row(t_ext, ext_headers)
for row_data in ext_data:
    add_data_row(t_ext, row_data, bold_first=True)
set_col_widths(t_ext, ext_widths)
doc2.add_paragraph()

# ── Processes ─────────────────────────────────────────────────────────────────
add_section_heading(doc2, '2. Processes', level=1)
add_caption(doc2, 'The seven functional processes that transform data within the system.')

proc_headers = ['ID', 'Name', 'Description', 'Key Inputs', 'Key Outputs']
proc_widths  = [0.35, 1.3, 2.1, 1.6, 1.45]

proc_data = [
    ('P1', 'User Authentication',
     'Handles registration, login, JWT issuance, email verification, and password reset.',
     'Credentials, reset token',
     'JWT token, user profile, verification email'),
    ('P2', 'Product Management',
     'CRUD operations for products and categories; Cloudinary image upload; stock tracking.',
     'Product data, category data, image files',
     'Product catalogue, image URLs, category tree'),
    ('P3', 'Order Processing',
     'Validates cart, reserves stock, creates order, manages status transitions (pending → delivered).',
     'Cart contents, shipping address, coupon code',
     'Order confirmation, stock update, payment intent'),
    ('P4', 'Payment Processing',
     'Initiates payment with the appropriate gateway, handles IPN / webhook callbacks, updates order and payment records.',
     'Payment request, gateway webhook payload',
     'Redirect URL, payment record, updated order status'),
    ('P5', 'Recommendation Engine',
     'Retrieves personalised or trending recommendations from the AI service; falls back to DB trending query if AI service is unavailable.',
     'User ID, product IDs from AI service',
     'Ranked and enriched product list delivered to customer'),
    ('P6', 'Review and Wishlist',
     'Creates, updates, and deletes product reviews; manages wishlist additions/removals; recalculates product rating aggregates.',
     'Review data, wishlist action (add/remove)',
     'Persisted review, updated product averageRating, wishlist state'),
    ('P7', 'Activity Tracking',
     'Logs user behavioural events (views, searches, cart actions, purchases) for analytics and AI model training data.',
     'Event type, product ID, metadata',
     'Activity log entry written to D8; aggregated data forwarded to AI service'),
]

t_proc = doc2.add_table(rows=1, cols=5)
style_table(t_proc)
add_header_row(t_proc, proc_headers)
for row_data in proc_data:
    add_data_row(t_proc, row_data, bold_first=True)
set_col_widths(t_proc, proc_widths)
doc2.add_paragraph()

# ── Data Stores ───────────────────────────────────────────────────────────────
add_section_heading(doc2, '3. Data Stores', level=1)
add_caption(doc2, 'Persistent storage locations used by the system processes.')

ds_headers = ['ID', 'Name', 'Technology', 'Description', 'Key Fields']
ds_widths  = [0.35, 1.0, 0.9, 2.2, 2.35]

ds_data = [
    ('D1', 'Users',        'MongoDB', 'All registered user accounts including credentials and embedded shipping addresses.',
     '_id, email, password (hashed), role, isVerified, addresses'),
    ('D2', 'Products',     'MongoDB', 'Product catalogue with pricing, inventory, imagery, and computed rating aggregates.',
     '_id, name, price, stock, status, category, images, averageRating'),
    ('D3', 'Orders',       'MongoDB', 'Customer orders from placement through fulfilment, including line-item snapshots.',
     '_id, user, items, status, paymentStatus, total, shippingAddress'),
    ('D4', 'Payments',     'MongoDB', 'Payment transaction records from all gateways with raw gateway response for audit.',
     '_id, order, method, status, transactionId, amount, currency, gatewayResponse'),
    ('D5', 'Cart',         'MongoDB', 'Persisted shopping carts — one per user — retained across sessions.',
     '_id, user, items, total, couponCode, discount'),
    ('D6', 'Reviews',      'MongoDB', 'Customer reviews and star ratings linked to products and users.',
     '_id, user, product, rating, body, verifiedPurchase, createdAt'),
    ('D7', 'Wishlist',     'MongoDB', 'Products saved by each user for future reference — one wishlist per user.',
     '_id, user, items[product, addedAt]'),
    ('D8', 'UserActivity', 'MongoDB', 'Raw behavioural event log: views, searches, cart actions, purchases, recommendation impressions.',
     '_id, user, product, activityType, metadata, timestamp'),
]

t_ds = doc2.add_table(rows=1, cols=5)
style_table(t_ds)
add_header_row(t_ds, ds_headers)
for row_data in ds_data:
    add_data_row(t_ds, row_data, bold_first=True)
set_col_widths(t_ds, ds_widths)
doc2.add_paragraph()

# ── Data Flows ────────────────────────────────────────────────────────────────
add_section_heading(doc2, '4. Data Flows', level=1)
add_caption(doc2, 'Named data flows between external entities, processes, and data stores.')

df_headers = ['Flow', 'From', 'To', 'Data', 'Description']
df_widths  = [0.45, 1.2, 1.2, 1.7, 2.25]

df_data = [
    ('F1',  'Customer (E1)',         'P1 Authentication',       'Email, password or registration form data',
     'User submits login credentials or new account registration'),
    ('F2',  'P1 Authentication',     'Customer (E1)',            'JWT access token, user profile object',
     'System returns auth token and profile on successful authentication'),
    ('F3',  'P1 Authentication',     'D1 Users',                'User document (create or update)',
     'New user created or existing user record updated in Users store'),
    ('F4',  'Admin (E2)',            'P2 Product Management',   'Product data, category data, image files',
     'Admin creates or modifies product listings and categories'),
    ('F5',  'P2 Product Management', 'D2 Products',             'Product document with Cloudinary image URLs',
     'Persist product data after Cloudinary upload completes'),
    ('F6',  'Customer (E1)',         'P3 Order Processing',     'Cart ID, shipping address, coupon code',
     'Customer initiates checkout'),
    ('F7',  'P3 Order Processing',   'D3 Orders',               'New order document with pending status',
     'Order written to store after cart validation and stock reservation'),
    ('F8',  'P3 Order Processing',   'P4 Payment Processing',   'Order ID, amount, payment method, phone / email',
     'Checkout triggers payment initiation'),
    ('F9',  'P4 Payment Processing', 'Payment Gateway (E3)',    'Payment request: amount, reference, redirect URLs',
     'System sends payment intent to selected gateway'),
    ('F10', 'Payment Gateway (E3)',  'P4 Payment Processing',   'IPN / webhook: status, transaction ID, raw payload',
     'Gateway notifies system of payment success or failure'),
    ('F11', 'P4 Payment Processing', 'D4 Payments',             'Payment record with final status',
     'Write payment outcome and raw gateway response for audit trail'),
    ('F12', 'P4 Payment Processing', 'D3 Orders',               'Updated paymentStatus and order status',
     'Order record updated following confirmed or failed payment'),
    ('F13', 'Customer (E1)',         'P7 Activity Tracking',    'Event type, product ID, metadata',
     'User interaction events (views, searches, cart actions) sent to tracker'),
    ('F14', 'P7 Activity Tracking',  'D8 UserActivity',         'Activity log entry',
     'Raw behavioural event persisted for analytics and AI training'),
    ('F15', 'P7 Activity Tracking',  'AI Service (E4)',         'Aggregated user behaviour summary',
     'Behaviour data forwarded to AI service for recommendation scoring'),
    ('F16', 'AI Service (E4)',       'P5 Recommendation Engine','Ranked product ID list with recommendation type',
     'AI service returns recommendation results for the requesting user'),
    ('F17', 'P5 Recommendation Engine','Customer (E1)',         'Enriched and ranked product list',
     'Personalised or trending products delivered to the customer UI'),
    ('F18', 'Customer (E1)',         'P6 Review and Wishlist',  'Review content (rating, body) or wishlist action',
     'User submits a review or adds/removes a wishlist item'),
    ('F19', 'P6 Review and Wishlist','D6 Reviews / D7 Wishlist','Review document or updated wishlist array',
     'Review or wishlist state persisted to respective data store'),
    ('F20', 'P6 Review and Wishlist','D2 Products',             'Updated averageRating and reviewCount',
     'Product aggregate rating fields recalculated after each new review'),
]

t_df = doc2.add_table(rows=1, cols=5)
style_table(t_df)
add_header_row(t_df, df_headers)
for row_data in df_data:
    add_data_row(t_df, row_data, bold_first=True)
set_col_widths(t_df, df_widths)

doc2.save(r'c:\Users\hp\Desktop\e_commerce\docs\DD_DFD.docx')
print('DD_DFD.docx saved')
